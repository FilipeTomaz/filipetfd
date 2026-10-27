"""Apresentação editável e PDF do cenário 65/12/60, sem alterar resultados anteriores.

Execute após Rscript scripts/calcular_pos_transicao.R, a partir de qualquer pasta.
Requer python-pptx, reportlab, matplotlib, numpy e fontes DejaVu Sans.
"""
from pathlib import Path
import csv
import hashlib
import json
import subprocess
import os
import tempfile

os.environ.setdefault('MPLCONFIGDIR', str(Path(tempfile.gettempdir()) / 'filipetfd-matplotlib'))
os.environ.setdefault('XDG_CACHE_HOME', str(Path(tempfile.gettempdir()) / 'filipetfd-cache'))

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import simpleSplit

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
DATA = ROOT / 'resultados/pos_transicao_65_12_60'
STEM = 'Beneficio_65_anos_12_pagamentos_60_elegiveis'
NAVY = '#142D45'
TEAL = '#008675'
GOLD = '#CB8F31'
BLUE = '#5B7C99'
GRAY = '#526575'
LIGHT = '#EDF3F5'
WHITE = '#FFFFFF'


def read_csv(filename):
    with (DATA / filename).open(encoding='utf-8') as f:
        return [{k: float(v) for k, v in row.items()} for row in csv.DictReader(f)]


annual = read_csv('anual.csv')
summary = read_csv('resumo_bilhoes.csv')
by_year = {int(r['ano']): r for r in annual}
summary_year = {int(r['ano']): r for r in summary}
assert len(annual) == 34 and len(summary) == 7
assert all(r['economia_elegibilidade_real'] == 0 for r in annual if r['ano'] < 2037)


def fmt(value, digits=3):
    return f'{value:,.{digits}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def bi(value):
    return fmt(value / 1e9)


# Figura científica exportável, com dados fornecidos pelos CSVs em R.
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 11,
                     'axes.spines.top': False, 'axes.spines.right': False})
years = np.array([int(r['ano']) for r in annual])
age = np.array([r['economia_idade_real'] / 1e9 for r in annual])
payments = np.array([r['economia_numero_real'] / 1e9 for r in annual])
eligibility = np.array([r['economia_elegibilidade_real'] / 1e9 for r in annual])
fig, ax = plt.subplots(figsize=(12.0, 3.6), layout='constrained')
ax.bar(years, age, color=NAVY, label='Elevação da idade')
ax.bar(years, payments, bottom=age, color=GOLD, label='13 para 12 pagamentos')
ax.bar(years, eligibility, bottom=age+payments, color=TEAL,
       label='60% elegíveis desde 2037')
ax.axvline(2036.5, color=GRAY, linestyle='--', linewidth=1)
ax.text(2036.8, 295, 'Início do filtro de elegibilidade', color=GRAY, fontsize=10)
ax.set_ylim(0, 320)
ax.set_xlim(2026.25, 2060.8)
ax.set_ylabel('Economia anual · R$ bilhões de 2026', color=GRAY)
ax.set_xticks([2027, 2030, 2036, 2040, 2050, 2060])
ax.tick_params(colors=GRAY)
ax.grid(axis='y', alpha=.15)
ax.set_axisbelow(True)
ax.legend(loc='upper left', frameon=False, ncol=1, fontsize=10)
for year in [2036, 2040, 2060]:
    value = by_year[year]['economia_total_real'] / 1e9
    ax.annotate(fmt(value, 1), (year, value), xytext=(0, 6),
                textcoords='offset points', ha='center', color=NAVY, fontsize=10)
chart = OUT / 'economia_anual.png'
fig.savefig(chart, dpi=240, facecolor='white')
fig.savefig(OUT / 'economia_anual.svg', facecolor='white')
plt.close(fig)

fontdir = Path('/usr/share/fonts/truetype/dejavu')
for name, filename in [('DejaVu', 'DejaVuSans.ttf'), ('DejaVu-Bold', 'DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(fontdir / filename)))
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
pdf = canvas.Canvas(str(OUT / f'{STEM}.pdf'), pagesize=(960, 540))
pdf.setTitle('Benefício: 65 anos, 12 pagamentos e 60% elegíveis após a transição')
pdf.setAuthor('Projeções a partir de FilipeTomaz/filipetfd')
slide = None
count = 0


def rect(x, y, w, h, color):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color[1:])
    shape.line.fill.background()
    pdf.setFillColor(color)
    pdf.rect(x*72, 540-(y+h)*72, w*72, h*72, fill=1, stroke=0)


def text(x, y, w, h, value, size=18, color=NAVY, bold=False, align='left'):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    font = 'DejaVu-Bold' if bold else 'DejaVu'
    baseline = 540-y*72-size
    for index, line in enumerate(str(value).split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = 'DejaVu Sans'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color[1:])
        paragraph.space_after = Pt(5)
        if align == 'right':
            paragraph.alignment = PP_ALIGN.RIGHT
        for wrapped in simpleSplit(line, font, size, w*72):
            pdf.setFont(font, size)
            pdf.setFillColor(color)
            if align == 'right':
                pdf.drawRightString((x+w)*72, baseline, wrapped)
            else:
                pdf.drawString(x*72, baseline, wrapped)
            baseline -= size*1.2
        baseline -= 5


def start(kicker, title, subtitle='', dark=False):
    global slide, count
    if count:
        pdf.showPage()
    count += 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(0, 0, 13.333333, 7.5, NAVY if dark else WHITE)
    rect(0, 0, 13.333333, .09, TEAL)
    text(.6, .35, 12.1, .35, kicker.upper(), 10, '#87D8C8' if dark else TEAL, True)
    text(.6, .88, 12.1, .68, title, 27, WHITE if dark else NAVY, True)
    if subtitle:
        text(.62, 1.66, 12, .6, subtitle, 14, '#C3D4E0' if dark else GRAY)
    rect(.6, 6.98, 12.1, .012, '#496376' if dark else '#D9E3E7')
    color = '#C3D4E0' if dark else GRAY
    text(.6, 7.12, 11.6, .22,
         'Fonte: FilipeTomaz/filipetfd · Planilha original · Valores reais de 2026 · Cenário hipotético',
         8, color)
    text(12.2, 7.09, .5, .25, f'{count}/8', 9, color, align='right')


def table(headers, rows, widths, y=2.35, row_h=.48, size=14, x=.65):
    shape = slide.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y),
                                   Inches(sum(widths)), Inches(row_h*(len(rows)+1)))
    tab = shape.table
    for j, width in enumerate(widths):
        tab.columns[j].width = Inches(width)
    for i, row in enumerate([headers]+rows):
        tab.rows[i].height = Inches(row_h)
        xpos = x
        for j, value in enumerate(row):
            last = j == len(widths)-1
            bg = NAVY if i == 0 else (TEAL if last else (LIGHT if i % 2 else WHITE))
            fg = WHITE if i == 0 or last else NAVY
            cell = tab.cell(i, j)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(bg[1:])
            cell.margin_left = cell.margin_right = Inches(.12)
            cell.margin_top = cell.margin_bottom = Inches(.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fs = size-1 if i == 0 else size
            bold = i == 0 or last
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
                p.font.name = 'DejaVu Sans'
                p.font.size = Pt(fs)
                p.font.bold = bold
                p.font.color.rgb = RGBColor.from_string(fg[1:])
            rect_pdf_y = 540-(y+(i+1)*row_h)*72
            pdf.setFillColor(bg)
            pdf.rect(xpos*72, rect_pdf_y, widths[j]*72, row_h*72, fill=1, stroke=0)
            pdf.setFillColor(fg)
            pdf.setFont('DejaVu-Bold' if bold else 'DejaVu', fs)
            baseline = 540-(y+(i+.5)*row_h)*72-fs*.35
            if j == 0:
                pdf.drawString((xpos+.12)*72, baseline, str(value))
            else:
                pdf.drawRightString((xpos+widths[j]-.12)*72, baseline, str(value))
            xpos += widths[j]


def bullet(y, title, detail, color=TEAL):
    rect(.67, y+.05, .07, .56, color)
    text(.95, y, 11.65, .35, title, 17, NAVY, True)
    text(.95, y+.43, 11.65, .65, detail, 13, GRAY)


last = by_year[2060]
start('01 · Resultado principal', '65 anos para homens e mulheres',
      'Transição escalonada de 10 anos · 12 pagamentos anuais · 60% elegíveis após a transição', dark=True)
text(.68, 2.55, 7.0, 1.45, f'R$ {fmt(last["economia_total_real_acumulada"]/1e12)}\ntrilhões', 37, WHITE, True)
text(.7, 4.35, 6.9, .8, 'Economia acumulada entre 2027 e 2060,\nem reais de 2026.', 19, '#C3D4E0')
rect(8.0, 2.45, 4.65, 3.65, '#203F57')
text(8.3, 2.77, 4.0, .35, 'COMO O CENÁRIO FUNCIONA', 11, '#87D8C8', True)
text(8.3, 3.4, 4.0, 1.1, '2027–2036\nElevação gradual até 65 anos', 19, WHITE, True)
text(8.3, 4.8, 4.0, 1.0, 'Desde 2037\n60% dos novos ingressantes', 19, WHITE, True)
text(.7, 6.24, 11.8, .43, 'Comparação com H60/M55, 13 pagamentos e sem filtro adicional. Coortes anteriores a 2027 preservadas.',
     12, '#C3D4E0')

start('02 · Regras da transição', 'Dez degraus anuais até chegar aos 65',
      'A exigência é atualizada em janeiro: +6 meses/ano para homens e +12 meses/ano para mulheres.')
rows = []
for year in range(2027, 2037):
    step = year-2026
    h = 60+.5*step
    rows.append([str(year), f'{int(h)} anos' + (' e 6 meses' if h % 1 else ''),
                 f'{55+step} anos'])
table(['Ano', 'Homens', 'Mulheres'], rows, [1.0, 3.7, 2.5], y=2.32, row_h=.34, size=12)
rect(8.25, 2.32, 4.4, 1.6, LIGHT)
text(8.5, 2.55, 3.9, .4, 'Durante a transição', 17, NAVY, True)
text(8.5, 3.08, 3.9, .7, '100% elegíveis pela regra etária.\n12 pagamentos desde 2027.', 15, GRAY)
rect(8.25, 4.2, 4.4, 1.86, TEAL)
text(8.5, 4.43, 3.9, .4, 'A partir de 2037', 17, WHITE, True)
text(8.5, 4.98, 3.9, .85, 'Idade: 65 anos para ambos.\n60% dos ingressos efetivos\npermanecem elegíveis.', 15, WHITE)
text(.68, 6.4, 11.9, .32, 'O corte de elegibilidade não é antecipado: nenhum efeito desse filtro é contabilizado antes de 2037.', 12, GRAY)

start('03 · Trajetória anual', 'A economia anual chega a R$ '+bi(last['economia_total_real'])+' bi em 2060',
      'Efeitos separados, sem dupla contagem · R$ bilhões de 2026 por ano')
slide.shapes.add_picture(str(chart), Inches(.6), Inches(2.24), width=Inches(12.1), height=Inches(3.63))
pdf.drawImage(str(chart), .6*72, 540-(2.24+3.63)*72, width=12.1*72, height=3.63*72)
text(.7, 6.15, 11.8, .5, 'A idade e os 12 pagamentos atuam desde 2027. O filtro de 60% acrescenta economia somente a partir de 2037.', 14, GRAY)

start('04 · Economia acumulada', 'Quanto cada mudança acrescenta à economia',
      'Acumulado desde 2027 · R$ bilhões de 2026 · A última coluna soma os três efeitos')
rows = []
for year in [2030, 2036, 2040, 2050, 2060]:
    r = by_year[year]
    rows.append([str(year), bi(r['economia_idade_real_acumulada']),
                 bi(r['economia_numero_real_acumulada']),
                 bi(r['economia_elegibilidade_real_acumulada']),
                 bi(r['economia_total_real_acumulada'])])
table(['Até o ano', 'Idade', '13 → 12', 'Elegibilidade', 'Total'], rows,
      [1.3, 2.7, 2.4, 2.8, 2.8], y=2.5, row_h=.56, size=17)
text(.7, 6.16, 11.9, .52, 'Decomposição sequencial: primeiro a idade; depois a retirada do 13º; por fim a elegibilidade com 12 pagamentos. Totais antes do arredondamento.', 12, GRAY)

start('05 · Primeiros anos', 'A transição e o primeiro ano do redutor',
      'Economia de cada ano · R$ bilhões de 2026 · O redutor começa em 2037')
rows = []
for year in range(2027, 2038):
    r = by_year[year]
    rows.append([str(year), bi(r['economia_idade_real']),
                 bi(r['economia_numero_real']), bi(r['economia_elegibilidade_real']),
                 bi(r['economia_total_real'])])
table(['Ano', 'Idade', '13 → 12', 'Elegibilidade', 'Total anual'], rows,
      [1.3, 2.7, 2.4, 2.8, 2.8], y=2.3, row_h=.335, size=12)
text(.7, 6.55, 11.9, .25, 'Em 2027, a economia decorre apenas dos 12 pagamentos: o modelo preserva a idade média efetiva da referência.', 10, GRAY)

start('06 · Despesa comparável', 'Referência e proposta no mesmo fluxo de coortes',
      'Despesa anual apenas das coortes de referência desde 2027 · R$ bilhões de 2026')
rows = []
for year in [2036, 2037, 2040, 2050, 2060]:
    r = by_year[year]
    rows.append([str(year), bi(r['despesa_referencia_real']),
                 bi(r['despesa_proposta_real']), bi(r['economia_total_real'])])
table(['Ano', 'Regra atual', 'Proposta', 'Economia anual'], rows,
      [1.5, 3.5, 3.5, 3.5], y=2.55, row_h=.55, size=17)
text(.75, 6.16, 11.7, .53, 'Estas despesas não representam todo o orçamento previdenciário. As coortes anteriores a 2027 são mantidas fora da comparação e não sofrem corte.', 13, GRAY)

start('07 · Aplicação do filtro', 'O ano do ingresso efetivo determina a elegibilidade',
      'O benefício mensal continua igual a um salário mínimo. O fator de 60% reduz o número de ingressantes.')
bullet(2.45, 'Ingresso efetivo até dezembro de 2036',
       'Não recebe o novo filtro de elegibilidade. O recebimento continua nos anos seguintes, com 12 pagamentos anuais.')
bullet(3.83, 'Ingresso efetivo em janeiro de 2037 ou depois',
       '60% permanecem elegíveis e 40% são excluídos deste benefício. A exclusão é permanente no exercício.')
bullet(5.21, 'Pessoa que alcançou a idade antiga antes de 2037, mas teve o ingresso adiado',
       'Se só ingressar pela nova idade a partir de 2037, passa pelo filtro. A data relevante é a concessão simulada.')

start('08 · Premissas e limites', 'Como interpretar os resultados',
      'Projeção condicional às premissas da planilha; não é uma avaliação individual de elegibilidade.')
bullet(2.35, 'Mesma base demográfica, econômica e histórica do projeto',
       'População, salário mínimo e INPC da planilha original; incidência média de concessões de 2020–2024.\nA população de referência H60/M55 não foi substituída por uma nova série demográfica aos 65 anos.')
bullet(3.85, 'Valores reais de 2026; acumulação sem desconto financeiro',
       'Pagamentos proporcionais aos meses de recebimento. O modelo mantém os benefícios após o ingresso,\nsem simular mortalidade, cessação ou migração para outros benefícios.')
bullet(5.35, 'O cenário incide somente sobre o fluxo modelado de novas coortes',
       'Coortes anteriores a 2027 preservadas. Os 40% excluídos não retornam a este benefício no horizonte.\nOs resultados não medem o gasto total do sistema nem efeitos comportamentais da mudança.')

prs.core_properties.title = '65 anos, 12 pagamentos e 60% elegíveis após transição de 10 anos'
prs.core_properties.subject = 'Economia de 2027 a 2060; valores reais de 2026'
prs.core_properties.author = 'Projeções a partir de FilipeTomaz/filipetfd'
prs.save(OUT / f'{STEM}.pptx')
pdf.save()
assert count == len(prs.slides) == 8
assert sum(shape.has_table for s in prs.slides for shape in s.shapes) == 4
metadata = {
    'repository': 'FilipeTomaz/filipetfd',
    'base_commit': subprocess.check_output(['git', '-C', str(ROOT), 'rev-parse', 'HEAD'], text=True).strip(),
    'workbook_sha256': hashlib.sha256((ROOT / 'Proposta de Equiparação da Idade pela DIB.xlsx').read_bytes()).hexdigest(),
    'annual_csv_sha256': hashlib.sha256((DATA / 'anual.csv').read_bytes()).hexdigest(),
    'payments_reference': 13, 'payments_proposal': 12,
    'transition_start': 2027, 'transition_end': 2036,
    'eligibility_start_actual_admission': 2037, 'eligible_share': .6,
    'real_currency_year': 2026, 'slides': 8,
    'cumulative_real_savings_2060': last['economia_total_real_acumulada'],
}
(OUT / 'metadados.json').write_text(json.dumps(metadata, indent=2, ensure_ascii=False)+'\n', encoding='utf-8')
print(f'Gerados: {STEM}.pptx, {STEM}.pdf e gráfico PNG/SVG; 8 slides, 4 tabelas.')
