"""Gera quatro slides em PPTX editável e PDF a partir dos resultados calculados em R.
Requer python-pptx e reportlab. Execute a partir de qualquer diretório.
"""
from pathlib import Path
import csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import simpleSplit

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
with (ROOT / 'resultados/tres_cenarios_resumo_bilhoes.csv').open(encoding='utf-8') as f:
    dados = list(csv.DictReader(f))
NAVY, GREEN, ORANGE, GRAY, LIGHT = '#142D45', '#008675', '#DC743B', '#526575', '#EDF3F5'
fontdir = Path('/usr/share/fonts/truetype/dejavu')
if not fontdir.exists():
    raise RuntimeError('Instale a fonte DejaVu Sans para gerar o PDF com a mesma tipografia.')
for name, filename in [('DejaVu', 'DejaVuSans.ttf'), ('DejaVu-Bold', 'DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(fontdir / filename)))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
pdf = canvas.Canvas(str(OUT / 'Comparacao_tres_cenarios.pdf'), pagesize=(960, 540))
pdf.setTitle('Aposentadoria rural — três cenários de transição')
slide = None

def rect(x, y, w, h, color):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color[1:])
    shape.line.fill.background()
    pdf.setFillColor(color)
    pdf.rect(x*72, 540-(y+h)*72, w*72, h*72, fill=1, stroke=0)

def text(x, y, w, h, value, size=18, color=NAVY, bold=False):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    font = 'DejaVu-Bold' if bold else 'DejaVu'
    pdf.setFillColor(color)
    pdf.setFont(font, size)
    yy = 540-y*72-size
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name, p.font.size, p.font.bold = 'DejaVu Sans', Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(color[1:])
        p.space_after = Pt(6)
        for wrapped in simpleSplit(line, font, size, w*72):
            pdf.drawString(x*72, yy, wrapped)
            yy -= size*1.2
        yy -= 6

def start(n, kicker, title):
    global slide
    if n > 1:
        pdf.showPage()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(0, 0, 13.333, .12, GREEN)
    text(.55, .35, 12, .3, kicker.upper(), 11, GREEN, True)
    text(.55, .86, 12.2, .9, title, 29, NAVY, True)
    rect(.55, 6.94, 12.2, .012, '#D9E3E7')
    text(.55, 7.08, 11.5, .25, 'Fonte: comparacao_tres_cenarios.Rmd • Valores reais: R$ de 2026 • Referência: regras atuais', 9, GRAY)
    text(12.2, 7.08, .5, .25, f'{n}/4', 9, GRAY)

def fmt(v):
    return f'{float(v):,.3f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

def table(headers, body, xs, widths, size=22):
    y0 = 2.45
    heights = [.72] + [.64]*len(body)
    tab = slide.shapes.add_table(len(body)+1, len(headers), Inches(xs[0]), Inches(y0),
                                Inches(sum(widths)), Inches(sum(heights))).table
    for j, width in enumerate(widths):
        tab.columns[j].width = Inches(width)
    for i, row in enumerate([headers]+body):
        tab.rows[i].height = Inches(heights[i])
        y = y0+sum(heights[:i])
        for j, value in enumerate(row):
            last = j == len(headers)-1
            bg = NAVY if i == 0 else (GREEN if last else (LIGHT if i % 2 else '#FFFFFF'))
            fg = '#FFFFFF' if i == 0 or last else NAVY
            fs = 17 if i == 0 else size
            cell = tab.cell(i, j)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(bg[1:])
            cell.margin_left = cell.margin_right = Inches(.18)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
                p.font.name, p.font.size, p.font.bold = 'DejaVu Sans', Pt(fs), (i == 0 or last)
                p.font.color.rgb = RGBColor.from_string(fg[1:])
            pdf.setFillColor(bg)
            pdf.rect(xs[j]*72, 540-(y+heights[i])*72, widths[j]*72, heights[i]*72, fill=1, stroke=0)
            pdf.setFillColor(fg)
            pdf.setFont('DejaVu-Bold' if i == 0 or last else 'DejaVu', fs)
            baseline = 540-(y+heights[i]/2)*72-fs*.35
            if j == 0:
                pdf.drawString((xs[j]+.18)*72, baseline, value)
            else:
                pdf.drawRightString((xs[j]+widths[j]-.18)*72, baseline, value)

start(1, '01 • Desenho do exercício', 'Três cenários, uma mesma referência')
text(.65, 1.85, 12, .8, 'Referência: regras atuais — homens aos 60 e mulheres aos 55 anos, sem restrição adicional de elegibilidade.', 21)
for i, (label, detail) in enumerate([
    ('1  |  Cenário base', 'Homens 65 • Mulheres 62 • 100% elegíveis'),
    ('2  |  Idades equiparadas', 'Homens 65 • Mulheres 65 • 100% elegíveis'),
    ('3  |  Elegibilidade restrita', 'Homens 65 • Mulheres 65 • 60% elegíveis')]):
    y = 2.95+i*.82
    rect(.65, y, 12.03, .65, LIGHT)
    text(.85, y+.16, 4.3, .38, label, 18, NAVY, True)
    text(5.2, y+.16, 7.1, .38, detail, 18)
text(.7, 5.66, 11.8, .7, 'Transição uniforme de 10 anos: 2027–2036.\nResultados acumulados desde 2027, em R$ bilhões de 2026.', 19, GREEN, True)
text(.7, 6.54, 11.8, .3, 'Referência de concessão pela idade média efetiva; preservação das coortes anteriores a 2027.', 11, GRAY)

for n, key, title, description in [
    (2, 'base_65_62', 'Cenário base: homens 65, mulheres 62', '100% elegíveis • transição de 10 anos • economia somente pelo aumento de idade'),
    (3, 'equiparado_65_65', 'Idades equiparadas: 65 anos para ambos', '100% elegíveis • transição de 10 anos • mesma referência do cenário base')]:
    start(n, f'0{n} • Resultados', title)
    text(.65, 1.83, 12, .6, description, 18, GRAY)
    values = [r for r in dados if r['exercicio'] == key]
    table(['Até o ano', 'Economia acumulada — R$ bi de 2026'],
          [[r['ano'], fmt(r['total_real_bi'])] for r in values],
          [1.3, 4.1], [2.8, 7.9], 24)
    text(1.35, 6.03, 10.6, .55, 'Sem restrição adicional de elegibilidade. Valores frente à manutenção das regras atuais.', 17)

start(4, '04 • Resultados', '65 anos para ambos, com 60% elegíveis')
text(.65, 1.83, 12, .45, 'Transição de 10 anos • economia acumulada em R$ bilhões de 2026', 19, GRAY)
values = [r for r in dados if r['exercicio'] == 'restrito_65_65_60']
table(['Até o ano', 'Aumento de idade', 'Exclusão dos 40%', 'Total'],
      [[r['ano'], fmt(r['idade_real_bi']), fmt(r['restricao_real_bi']), fmt(r['total_real_bi'])] for r in values],
      [.65, 2.25, 6.0, 9.75], [1.6, 3.75, 3.75, 2.3], 21)
text(.7, 5.96, 11.9, .65, 'Total = adiamento entre os 60% elegíveis + pagamentos evitados pela exclusão dos demais 40%.', 17, NAVY, True)
text(.7, 6.67, 11.9, .22, 'Por hipótese, os excluídos não recebem este benefício posteriormente. Componentes arredondados separadamente.', 10, GRAY)
prs.core_properties.title = 'Aposentadoria rural — três cenários de transição'
prs.core_properties.subject = 'Tabelas de economia acumulada até 2060'
prs.save(OUT / 'Comparacao_tres_cenarios.pptx')
pdf.save()
assert len(prs.slides) == 4
assert sum(sh.has_table for s in prs.slides for sh in s.shapes) == 3
print('PPTX e PDF gerados: quatro slides, com três tabelas editáveis no PowerPoint.')
